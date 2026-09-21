"""Prueft die statischen Folien der Themen-Vorlagen (21.09.2026).

AUSLOESER: Philip hat die aktuellen Original-Broschueren aus dem Haus
geliefert (Pro, Offensiv, Pro Dividende, Stand 14.09.2026). Die taeglich
befuellten Folien passten, die STATISCHEN wichen ab — gemessen am Folientext
und am PowerPoint-Bild:

    alle          "Unsere Bank in Zahlen" noch mit den Zahlen 31.12.2024
                  (comdirect ebenso, Nachtrag am selben Tag)
    Pro, Pro Div. "Steuerlicher Hinweis zum Honorar" fehlte
    Offensiv      eigenes Cover, eigene Leitlinien, eigene Honorar-Folie,
                  KEINE Folie "Gute Jahre ueberwiegen" (20 statt 22 Folien)
    Pro Div.      Honorar-Tabelle zeigte "Strategie Pro"
    Offensiv      dito — dort auch im Original, auf Zuruf korrigiert

Die Folien wurden 1:1 aus den Originalen in die Vorlagen uebernommen. Beide
SCHWEIZ-Strategien nutzen die Pro-Vorlage und bekommen die Aenderungen mit.
Nichts im Code erzeugt diese Folien — eine Vorlage, die jemand spaeter aus
einem alten Stand zurueckholt, faellt nur hier auf.

Dazu das Datum der Schlussfolie: "Stand: …" ist ein PowerPoint-DATUMSFELD.
PowerPoint zeigt beim Oeffnen das heutige Datum; der gespeicherte Wert (den
Vorschauen ohne Feldaktualisierung zeigen) stand auf 06.07.2026 und wird
jetzt beim Export auf den Datenstand gesetzt (`update_stand_datum`).

    Schritt 1 — Folienfolge je Vorlage namentlich
    Schritt 2 — Kernwerte der uebernommenen Folien
    Schritt 3 — Offensiv hat ein anderes Cover als Pro
    Schritt 4 — keine Personennamen in den Metadaten (auch eingebettete Excel)
    Schritt 5 — update_stand_datum setzt den Feldwert, das Feld bleibt
    Schritt 6 — Datums-Stimmigkeit ALLER Vorlagen (Nachtrag 21.09.2026):
                Bank-Folie in sich stimmig (Balkenkopf = Kennzahlen = letzter
                Balken), kein fester "Stand: TT.MM.JJJJ" als Folientext
    Schritt 7 — Gegenproben: die alten Zustaende wuerden gemeldet

    python tests/test_thema_statische_folien.py
"""

import hashlib
import io
import os
import re
import sys
import zipfile

WURZEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, WURZEL)
FIRMA = "Fürst Fugger Privatbank"

_ANFANG = ["Unsere Vermögensverwaltung"]
_MITTE = [
    "Aktien – die guten Jahre überwiegen",
    "Die Fallstricke des typischen Investors",
    "Durchhalten zahlt sich aus",
]
_BLOCK = [
    "Einzeltitel",
    "Aktuelle Portfoliozusammenstellung",
    "Anlagestrategie",                  # Titel wird beim Export befuellt
    "Wertentwicklung der Strategie",    # dito
]
_ENDE = [
    "Unsere Bank in Zahlen",
    "Unsere Standorte",
    "Unsere Standorte",
    "Risikohinweise",
    "Rechtliche Hinweise und Impressum",
    "Vielen Dank für Ihr Interesse",
]


def _pro_folge(strategie_titel):
    return (_ANFANG + [strategie_titel] * 2 + _MITTE
            + ["Gute Jahre überwiegen", "Krise als Chance",
               "Basis unserer Investmententscheidungen"]
            + _BLOCK
            + ["Unser Honorar", "Unser Honorar", "Steuerlicher Hinweis zum Honorar"]
            + _ENDE)


# Vorlage -> erwartete Titel (Anfang des Titels, Folie fuer Folie)
FOLGE = {
    "Vorlage_Thema.pptx": _pro_folge("Unsere Strategie PRO"),
    "Vorlage_Thema_ProDividende.pptx": _pro_folge("Unsere Strategie Pro Dividende"),
    "Vorlage_Thema_Offensiv.pptx": (
        _ANFANG + ["Unsere Strategie „Offensiv“"] * 2 + _MITTE
        + ["Krise als Chance", "Basis unserer Investmententscheidungen"]
        + _BLOCK
        + ["Unser Honorar", "Unser Honorar"]
        + _ENDE),
}

# Vorlage -> [(Folientitel, Text, der auf dieser Folie stehen MUSS)]
KERNWERTE = {
    "Vorlage_Thema.pptx": [
        ("Unsere Bank in Zahlen", "Dezember 2025"),
        ("Unsere Bank in Zahlen", "7.358 Mio. EUR"),
        ("Steuerlicher Hinweis zum Honorar", "Abgeltungssteuer"),
    ],
    "Vorlage_Thema_ProDividende.pptx": [
        ("Unsere Bank in Zahlen", "Dezember 2025"),
        ("Unsere Bank in Zahlen", "7.358 Mio. EUR"),
        ("Steuerlicher Hinweis zum Honorar", "Abgeltungssteuer"),
        ("Unser Honorar", "Strategie Pro Dividende"),
    ],
    "Vorlage_Thema_Offensiv.pptx": [
        ("Unsere Bank in Zahlen", "Dezember 2025"),
        ("Unsere Bank in Zahlen", "7.358 Mio. EUR"),
        ("Basis unserer Investmententscheidungen", "Aktienrückkäufe"),
        ("Unser Honorar", "Halbjährliche Abrechnung"),
        # Das Original traegt hier "Strategie Pro" — auf Zuruf Philip
        # korrigiert (21.09.2026), weicht also bewusst vom Original ab.
        ("Unser Honorar", "Strategie Offensiv"),
    ],
    # Nachtrag 21.09.2026: comdirect F21 trug noch Dezember 2024; die Folie
    # ist jetzt die (layoutgleiche) aktuelle Bank-Folie der cVV-Vorlage.
    "Vorlage_comdirect.pptx": [
        ("Unsere Bank in Zahlen", "Dezember 2025"),
        ("Unsere Bank in Zahlen", "7.358 Mio. EUR"),
    ],
}


def _norm(text):
    return re.sub(r"\s+", " ", text.replace("\x0b", " ")).strip().casefold()


def _titel(folie):
    """Titel der Folie: Titel-Platzhalter zuerst, sonst der oberste Text."""
    kandidaten = []
    for sh in folie.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            ist_titel = sh.is_placeholder and str(
                sh.placeholder_format.type).startswith(("TITLE", "CENTER_TITLE"))
            kandidaten.append((0 if ist_titel else 1, sh.top, sh.text_frame.text))
    kandidaten.sort(key=lambda k: (k[0], k[1]))
    return kandidaten[0][2] if kandidaten else ""


def _texte(folie):
    """Aller Text einer Folie, auch aus Gruppen und Tabellen."""
    teile = []

    def sammeln(shapes):
        for sh in shapes:
            if sh.shape_type == 6:                     # Gruppe
                sammeln(sh.shapes)
            if sh.has_text_frame:
                teile.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for zeile in sh.table.rows:
                    teile.extend(z.text for z in zeile.cells)
    sammeln(folie.shapes)
    return _norm(" ".join(teile))


def pruefe_folge(prs, erwartet):
    titel = [_titel(f) for f in prs.slides]
    if len(titel) != len(erwartet):
        return [f"{len(titel)} Folien statt {len(erwartet)}"]
    return [f"Folie {i}: {t.strip()[:40]!r} statt {e!r}"
            for i, (t, e) in enumerate(zip(titel, erwartet), 1)
            if not _norm(t).startswith(_norm(e))]


def pruefe_kernwerte(prs, erwartet):
    fehler = []
    for titel, text in erwartet:
        folien = [f for f in prs.slides if _norm(_titel(f)).startswith(_norm(titel))]
        if not any(_norm(text) in _texte(f) for f in folien):
            fehler.append(f"{text!r} steht auf keiner Folie {titel!r}")
    return fehler


def _cover_bilder(prs):
    """Bild-Hashes der Titelfolie samt Layout."""
    folie = prs.slides[0]
    hashes = set()
    for teil in (folie.part, folie.slide_layout.part):
        for rel in teil.rels.values():
            if "image" in rel.reltype and not rel.is_external:
                hashes.add(hashlib.md5(rel.target_part.blob).hexdigest())
    return hashes


def autor_funde(pfad):
    """Autorfelder, die nicht der Firmenname sind — auch in Einbettungen."""
    funde = []

    def scan(name, blob):
        z = zipfile.ZipFile(io.BytesIO(blob))
        for n in z.namelist():
            if n.endswith("core.xml"):
                xml = z.read(n).decode("utf-8", "replace")
                for tag in ("dc:creator", "cp:lastModifiedBy"):
                    for wert in re.findall(rf"<{tag}>(.*?)</{tag}>", xml, re.S):
                        if wert.strip() and wert.strip() != FIRMA:
                            funde.append(f"{name}>{n} {tag}={wert!r}")
            elif n.endswith((".xlsx", ".xlsm", ".docx", ".pptx")):
                scan(f"{name}>{n}", z.read(n))
    with open(pfad, "rb") as f:
        scan(os.path.basename(pfad), f.read())
    return funde


def _feldwerte(prs):
    """Gespeicherte Werte aller Datumsfelder auf allen Folien."""
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    return [t.text for folie in prs.slides
            for fld in folie._element.xpath(".//a:fld[starts-with(@type,'datetime')]")
            for t in fld.findall("a:t", ns)]


_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_STAND_FEST = re.compile(r"(?i)\bstand[:.,]?\s*\d{2}\.\d{2}\.\d{4}")


def feste_stand_daten(prs):
    """'Stand: TT.MM.JJJJ' als FESTER Folientext (nicht in einem Feld).

    AUSLOESER: Das ESG-Impressum trug "Stand: 31.07.2024" als Text — in
    jeder gebauten ESG-Broschuere, zwei Jahre alt. Alle anderen Impressen
    haben dort ein Datumsfeld. Die Quellzeilen der Diagramme sind Felder
    (TxLink) und werden vom Export gesetzt; sie zaehlen hier nicht.
    """
    funde = []
    for i, folie in enumerate(prs.slides, 1):
        for p in folie._element.iter(_A + "p"):
            text = "".join(t.text or "" for t in p.iter(_A + "t")
                           if t.getparent().tag != _A + "fld")
            funde += [f"F{i}: {m.group(0)!r}" for m in _STAND_FEST.finditer(text)]
    return funde


def bank_jahre(prs):
    """Jahre auf der Folie 'Unsere Bank in Zahlen': (Kopf, Kennzahlen, letzter Balken).

    AUSLOESER: Nach dem Nachziehen auf 2025 stand im Balkenkopf weiter
    "DEZEMBER 2024" — auch im Original des Hauses. Drei Angaben derselben
    Folie, die nur gemeinsam stimmen.
    """
    for folie in prs.slides:
        if not _norm(_titel(folie)).startswith(_norm("Unsere Bank in Zahlen")):
            continue
        kennzahlen = re.search(r"auf einen blick\W*(?:dezember\s+|31\.12\.)(20\d\d)", _texte(folie))
        kopf = balken = None
        for sh in folie.shapes:
            if not getattr(sh, "has_chart", False) or not sh.has_chart:
                continue
            kategorien = list(sh.chart.plots[0].categories)
            balken = str(kategorien[-1]).strip()[:4] if kategorien else None
            for rel in sh.chart.part.rels.values():
                if rel.reltype.endswith("/chartUserShapes"):
                    txt = " ".join(t.text or "" for t in
                                   __import__("lxml.etree").etree.fromstring(
                                       rel.target_part.blob).iter(_A + "t"))
                    m = re.search(r"(?i)assets under control\s+(?:dezember\s+|31\.12\.)(20\d\d)", txt)
                    kopf = m.group(1) if m else None
        return (kopf, kennzahlen.group(1) if kennzahlen else None, balken)
    return None


def main():
    from importlib.util import find_spec
    if find_spec("pptx") is None:
        print("UEBERSPRUNGEN — python-pptx nicht installiert")
        return 0
    from pptx import Presentation
    from modules.pptx_helpers import update_stand_datum

    def laden(name):
        return Presentation(os.path.join(WURZEL, "Vorlage", name))

    f = 0
    print("Schritt 1 — Folienfolge je Vorlage")
    for vorlage, folge in FOLGE.items():
        abw = pruefe_folge(laden(vorlage), folge)
        for a in abw:
            print(f"   FEHLER — {vorlage}: {a}")
        f += len(abw)
        if not abw:
            print(f"   OK — {vorlage}: {len(folge)} Folien in der Reihenfolge des Originals")

    print("Schritt 2 — Kernwerte der uebernommenen Folien")
    for vorlage, werte in KERNWERTE.items():
        abw = pruefe_kernwerte(laden(vorlage), werte)
        for a in abw:
            print(f"   FEHLER — {vorlage}: {a}")
        f += len(abw)
        if not abw:
            print(f"   OK — {vorlage}: {len(werte)} Kernwerte gefunden")
    pro = laden("Vorlage_Thema.pptx")
    for vorlage, fremd in (("Vorlage_Thema.pptx", "Strategie Pro Dividende"),
                           ("Vorlage_Thema.pptx", "Strategie Offensiv")):
        if pruefe_kernwerte(laden(vorlage), [("Unser Honorar", fremd)]):
            print(f"   OK — {vorlage} traegt in der Honorar-Tabelle nicht {fremd!r}")
        else:
            print(f"   FEHLER — {vorlage} zeigt {fremd!r}")
            f += 1

    print("Schritt 3 — Offensiv hat sein eigenes Cover")
    if _cover_bilder(laden("Vorlage_Thema_Offensiv.pptx")) & _cover_bilder(pro):
        print("   FEHLER — Offensiv teilt Titelbilder mit Pro (altes Kommoden-Cover?)")
        f += 1
    else:
        print("   OK — keine gemeinsamen Titelbilder mit Pro")

    print("Schritt 4 — Metadaten ohne Personennamen")
    for vorlage in sorted(set(FOLGE) | set(KERNWERTE)):
        funde = autor_funde(os.path.join(WURZEL, "Vorlage", vorlage))
        for x in funde:
            print(f"   FEHLER — {x}")
        f += len(funde)
        if not funde:
            print(f"   OK — {vorlage}")

    print("Schritt 5 — Datumsfelder (Schlussfolie, Impressum)")
    datum = "18.09.2026"
    for vorlage in sorted(os.listdir(os.path.join(WURZEL, "Vorlage"))):
        if not vorlage.endswith(".pptx"):
            continue
        prs = laden(vorlage)
        vorher = _feldwerte(prs)
        update_stand_datum(prs, datum)
        nachher = _feldwerte(prs)
        if not vorher:
            print(f"   OK — {vorlage}: kein Datumsfeld, nichts zu tun")
        elif nachher == [datum] * len(vorher):
            print(f"   OK — {vorlage}: {vorher} -> {nachher} (Feld erhalten)")
        else:
            print(f"   FEHLER — {vorlage}: {vorher} -> {nachher}")
            f += 1

    print("Schritt 6 — Datums-Stimmigkeit aller Vorlagen")
    for vorlage in sorted(os.listdir(os.path.join(WURZEL, "Vorlage"))):
        if not vorlage.endswith(".pptx"):
            continue
        prs = laden(vorlage)
        fest = feste_stand_daten(prs)
        for x in fest:
            print(f"   FEHLER — {vorlage}: fester Stand-Text {x} (Datumsfeld verwenden)")
        f += len(fest)
        jahre = bank_jahre(prs)
        if jahre is None:
            print(f"   OK — {vorlage}: kein fester Stand-Text, keine Bank-Folie")
        elif None in jahre or len(set(jahre)) != 1:
            print(f"   FEHLER — {vorlage}: Bank-Folie widerspruechlich "
                  f"(Kopf {jahre[0]}, Kennzahlen {jahre[1]}, letzter Balken {jahre[2]})")
            f += 1
        else:
            print(f"   OK — {vorlage}: kein fester Stand-Text, Bank-Folie durchgehend {jahre[0]}")

    print("Schritt 7 — Gegenproben")
    # (a) Die Folie "Steuerlicher Hinweis" herausnehmen -> Schritt 1 meldet es.
    prs = laden("Vorlage_Thema.pptx")
    lst = prs.slides._sldIdLst
    lst.remove(lst[15])
    if pruefe_folge(prs, FOLGE["Vorlage_Thema.pptx"]):
        print("   OK — eine fehlende Steuer-Folie wuerde gemeldet")
    else:
        print("   FEHLER — fehlende Steuer-Folie ginge durch")
        f += 1
    # (b) Die alten Bankzahlen -> Schritt 2 meldet es.
    prs = laden("Vorlage_Thema.pptx")
    for folie in prs.slides:
        for sh in folie.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        r.text = r.text.replace("Dezember 2025", "31.12.2024")
    if pruefe_kernwerte(prs, KERNWERTE["Vorlage_Thema.pptx"]):
        print("   OK — die alten Bankzahlen wuerden gemeldet")
    else:
        print("   FEHLER — die alten Bankzahlen gingen durch")
        f += 1
    # (c) Ohne update_stand_datum bleibt der Vorlagenwert stehen.
    prs = laden("Vorlage_Thema.pptx")
    if _feldwerte(prs) != [datum]:
        print(f"   OK — ohne Aufruf steht {_feldwerte(prs)}, nicht der Datenstand")
    else:
        print("   FEHLER — Gegenprobe leer: Vorlage traegt schon den Datenstand")
        f += 1
    # (d) Der alte ESG-Impressumstext wuerde als fester Stand-Text gemeldet.
    prs = laden("Vorlage_Thema.pptx")
    prs.slides[-1].shapes.add_textbox(0, 0, 100, 100).text_frame.text = "Stand: 31.07.2024"
    if feste_stand_daten(prs):
        print("   OK — ein fester 'Stand: 31.07.2024' wuerde gemeldet")
    else:
        print("   FEHLER — fester Stand-Text ginge durch")
        f += 1
    # (e) Der alte Balkenkopf "DEZEMBER 2024" neben 2025er Zahlen faellt auf.
    prs = laden("Vorlage_Thema.pptx")
    for folie in prs.slides:
        for sh in folie.shapes:
            if getattr(sh, "has_chart", False) and sh.has_chart:
                for rel in sh.chart.part.rels.values():
                    if rel.reltype.endswith("/chartUserShapes"):
                        rel.target_part._blob = rel.target_part.blob.replace(
                            "DEZEMBER 2025".encode(), "DEZEMBER 2024".encode())
    j = bank_jahre(prs)
    if j and len(set(j)) != 1:
        print(f"   OK — der alte Balkenkopf wuerde gemeldet {j}")
    else:
        print(f"   FEHLER — alter Balkenkopf ginge durch {j}")
        f += 1

    print()
    if f:
        print(f"FEHLGESCHLAGEN — {f} Abweichung(en)")
        return 1
    print("BESTANDEN — statische Themen-Folien entsprechen den Originalen vom 14.09.2026")
    return 0


if __name__ == "__main__":
    sys.exit(main())
