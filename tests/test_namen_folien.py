"""Strategienamen auf den Folien — stimmig in sich und laut bei Umbenennung.

HINTERGRUND (24.09.2026, Etappe 5 der Stammdaten-Arbeit):
    Etappe 1 hat gefragt, ob eine Umbenennung im Mapping in der Broschuere
    ankommt. Antwort: nur halb — die Titel der Struktur-Folien stehen FEST in
    den Infoboard-Vorlagen, der Export fasst sie nie an (titel_text="").

    Beim Nachmessen fiel auf, dass es nicht erst bei einer Umbenennung
    knirscht, sondern schon heute: Die Wertentwicklungs-Folie direkt hinter
    der Struktur-Folie bekam ihren Titel aus dem MAPPING. Dieselbe Strategie
    hiess so auf zwei Folien hintereinander verschieden:

        cVV F11/12       Defensiv Plus        / Defensiv plus
        ESG F16-23       ESG Defensiv Plus    / ESG defensiv+
        ETF F16-19       ESG-ETF Ausgewogen   / ETF ausgewogen
        comdirect F6-11  Portfolioverwaltung 30 / Comdirect 30
                         (+ Kasten: "Anlagekriterien | FFPB Strategie 30")

    Entscheidung Philip: Die Wertentwicklungs-Folie uebernimmt den Namen von
    der Struktur-Folie (pptx_slides.strategiename_aus_titel), comdirect heisst
    ueberall "Portfolioverwaltung 30/70/100".

KOCHREZEPT — eine Strategie umbenennen (z. B. "Defensiv Plus" -> "X"):
    1. Mapping_Strategien.xlsx, Spalte A (Anzeigename im Tool) und Spalte
       "Anzeigename" (Kopfzeile des Kriterien-Kastens).
    2. Code-Konstanten, die den Mapping-Namen fuehren — tests/test_stammdaten.py
       Schritt 6 nennt jede einzelne.
    3. In der PowerPoint-Vorlage VON HAND:
         - Titel der Struktur-Folie "Anlagestrategie X" (die
           Wertentwicklungs-Folie folgt ihm von selbst)
         - Spaltenkopf der Uebersichtstabelle (cVV F17, ESG F24, ETF F20)
         - cVV F19: Serienname im Vergleichs-Chart (Daten bearbeiten)
    4. Code: VERGLEICH_FARBEN (pptx_export.py) bei cVV,
       _STRATEGIE_FAMILIE (chart_dynamik.py) mit dem neuen Titelnamen.
    Danach diesen Test laufen lassen — er meldet jede Stelle, die noch fehlt.

Geprueft wird:
  1. Jede Struktur-Folie der Infoboard-Vorlagen hat einen Titel
     "Anlagestrategie <Name>" — sonst faellt die Wertentwicklungs-Folie
     still auf den Mapping-Namen zurueck.
  2. Kasten-Name (Mapping, Spalte "Anzeigename") passt zum Titel: gleich
     oder dessen Ende ("Ausgewogen" zu "ESG-ETF Ausgewogen").
  3. Uebersichtstabelle: Spaltenkoepfe in Strategie-Reihenfolge, jeder
     beginnt mit dem Kasten-Namen ("Konservative Strategie" / "Konservativ").
  4. cVV-Vergleichs-Chart: Seriennamen = Kasten-Namen, in Reihenfolge.
  5. Titel -> Familie (_STRATEGIE_FAMILIE): jede Strategie wird ueber den
     Titel, den sie in der Broschuere traegt, ihrer Mapping-Familie
     zugeordnet; kein Eintrag der Tabelle ist tot.
  6. Gebaute Broschueren: Wertentwicklungs-Titel = Titelname der
     Struktur-Folie davor, und die Familie wird am Artefakt erkannt.
  7. Gegenproben — jeder der Schritte 1 bis 6 wird absichtlich
     verletzt und muss anschlagen.

    python tests/test_namen_folien.py [ausgabeordner]

Schritte 1-5 brauchen python-pptx und pandas. Schritt 6 baut echte
Broschueren und braucht zusaetzlich streamlit und die Daten; fehlt etwas,
wird er uebersprungen statt zu scheitern.
"""

import contextlib
import io
import os
import re
import sys
import tempfile

WURZEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, WURZEL)
sys.path.insert(0, os.path.join(WURZEL, "tests"))
os.chdir(WURZEL)

try:
    import pandas as pd
    from pptx import Presentation
except ImportError as ex:
    print(f"UEBERSPRUNGEN — {ex}")
    sys.exit(0)

from modules import stammdaten as stamm                          # noqa: E402
from modules import vorlagen_config as vc                        # noqa: E402
from modules import chart_dynamik as cd                          # noqa: E402
from modules.pptx_slides import (                                # noqa: E402
    WE_TITLE_FORMAT, clean_strategy_name, strategiename_aus_titel,
)

# Die cVV-Vergleichsfolie ist die einzige Einmal-Folie mit Strategienamen
# als Chart-Serien. Position kommt aus der Config, nicht als Zahl.
VERGLEICH_FAMILIE = "CVV"


def _vorlage(familie):
    datei, cfg = vc.VORLAGEN_FAMILIEN[familie]
    return Presentation(os.path.join("Vorlage", datei)), cfg


def _kasten_namen(sd):
    """Anzeigename im Tool -> Name im Kriterien-Kasten (Spalte 'Anzeigename')."""
    sp_a = stamm.spalte(sd, stamm.SP_ANZEIGE)
    sp_n = stamm.spalte(sd, stamm.SP_ANZEIGENAME)
    raus = {}
    for _, r in sd.iterrows():
        wert = r[sp_n]
        raus[str(r[sp_a]).strip()] = (None if pd.isna(wert) or not str(wert).strip()
                                      else str(wert).strip())
    return raus


def _titelnamen(prs, cfg, strategien):
    """(Strategie, Folie, Titelname) je Struktur-Folie der Vorlage."""
    raus = []
    for k, s in enumerate(strategien):
        pos = cfg["feste_bloecke"][k].get("anlagevorschlag")
        if pos is None:
            continue
        raus.append((s, pos, strategiename_aus_titel(prs.slides[pos - 1])))
    return raus


def _flach(text):
    """Zeilen- und Mehrfach-Leerzeichen zu einem Leerzeichen."""
    return " ".join(str(text).replace("\x0b", " ").split())


# ─────────────────────────────────────────────────────────────────────────
# 1. Struktur-Folien tragen "Anlagestrategie <Name>"
# ─────────────────────────────────────────────────────────────────────────
def _pruefe_titel_vorhanden(vorlagen):
    print("1. Struktur-Folien: Titel 'Anlagestrategie <Name>'")
    fehler = n = 0
    for fam, (prs, cfg) in vorlagen.items():
        for s, pos, name in _titelnamen(prs, cfg, vc.FAMILIE_ALLE_STRATEGIEN[fam]):
            n += 1
            if name is None:
                fehler += 1
                print(f"   FEHLER — {fam} F{pos} ({s}): Titel beginnt nicht mit "
                      f"'Anlagestrategie '. Die Wertentwicklungs-Folie faellt "
                      f"dann still auf den Mapping-Namen zurueck.")
    if not fehler:
        print(f"   OK — {n} Struktur-Folien in {len(vorlagen)} Vorlagen")
    return fehler


# ─────────────────────────────────────────────────────────────────────────
# 2. Kasten-Name <-> Titel
# ─────────────────────────────────────────────────────────────────────────
def _pruefe_kasten(vorlagen, kasten):
    print("\n2. Kriterien-Kasten (Mapping 'Anzeigename') passt zum Folientitel")
    fehler = n = 0
    for fam, (prs, cfg) in vorlagen.items():
        for s, pos, name in _titelnamen(prs, cfg, vc.FAMILIE_ALLE_STRATEGIEN[fam]):
            k = kasten.get(s)
            if name is None or k is None:
                continue          # Schritt 1 bzw. test_stammdaten Schritt 5
            n += 1
            if not (name == k or name.endswith(" " + k)):
                fehler += 1
                print(f"   FEHLER — {fam} F{pos} ({s}): Titel '{name}', Kasten "
                      f"'{k}'. Umbenannt? Kochrezept im Kopf dieser Datei.")
    if not fehler:
        print(f"   OK — {n} Kaesten, jeder gleich dem Titel oder dessen Ende")
    return fehler


# ─────────────────────────────────────────────────────────────────────────
# 3. Uebersichtstabelle
# ─────────────────────────────────────────────────────────────────────────
def _uebersicht_koepfe(prs, cfg):
    pos = cfg.get("einmal_folien", {}).get("uebersicht")
    if pos is None:
        return None, None
    for sh in prs.slides[pos - 1].shapes:
        if getattr(sh, "has_table", False):
            zellen = [_flach(c.text) for c in sh.table.rows[0].cells]
            if zellen and zellen[0].lower() == "zeitraum":
                return pos, [z for z in zellen[1:] if z]
    return pos, []


def _pruefe_uebersicht(vorlagen, kasten):
    print("\n3. Uebersichtstabelle: Spaltenkoepfe in Strategie-Reihenfolge")
    fehler = n = 0
    for fam, (prs, cfg) in vorlagen.items():
        pos, koepfe = _uebersicht_koepfe(prs, cfg)
        if pos is None:
            continue              # comdirect hat keine Uebersicht
        soll = [kasten.get(s) for s in vc.FAMILIE_ALLE_STRATEGIEN[fam]]
        n += 1
        if len(koepfe) != len(soll):
            fehler += 1
            print(f"   FEHLER — {fam} F{pos}: {len(koepfe)} Spaltenkoepfe, "
                  f"{len(soll)} Strategien")
            continue
        for kopf, k in zip(koepfe, soll):
            if k is None or not kopf.lower().startswith(k.lower()):
                fehler += 1
                print(f"   FEHLER — {fam} F{pos}: Spaltenkopf '{kopf}' passt "
                      f"nicht zu '{k}'")
    if not fehler:
        print(f"   OK — {n} Tabellen, jeder Kopf beginnt mit dem Kasten-Namen")
    return fehler


# ─────────────────────────────────────────────────────────────────────────
# 4. cVV-Vergleichs-Chart
# ─────────────────────────────────────────────────────────────────────────
def _pruefe_vergleich(vorlagen, kasten):
    print("\n4. cVV-Vergleichs-Chart: Seriennamen = Kasten-Namen")
    prs, cfg = vorlagen[VERGLEICH_FAMILIE]
    pos = cfg.get("einmal_folien", {}).get("vergleich")
    serien = None
    for sh in prs.slides[pos - 1].shapes:
        if getattr(sh, "has_chart", False):
            serien = [s.name for s in sh.chart.plots[0].series]
    soll = [kasten.get(s) for s in vc.FAMILIE_ALLE_STRATEGIEN[VERGLEICH_FAMILIE]]
    if serien != soll:
        print(f"   FEHLER — F{pos}: Serien {serien}, erwartet {soll}. "
              f"Serien in PowerPoint umbenennen (Daten bearbeiten), dann "
              f"VERGLEICH_FARBEN nachziehen (test_farben.py Schritt 5).")
        return 1
    print(f"   OK — F{pos}: {len(serien)} Serien in Strategie-Reihenfolge")
    return 0


# ─────────────────────────────────────────────────────────────────────────
# 5. Titel -> Familie
# ─────────────────────────────────────────────────────────────────────────
def _titel_je_strategie(vorlagen, sd):
    """Anzeigename -> (Familie aus dem Mapping, Titelname in der Broschuere).

    Infoboard-Familien: der Titel der Struktur-Folie. Alle anderen: der
    bereinigte Mapping-Name, so wie fill_wertentwicklung_slide ihn setzt.
    """
    sp_a = stamm.spalte(sd, stamm.SP_ANZEIGE)
    sp_f = stamm.spalte(sd, stamm.SP_FAMILIE)
    aus_vorlage = {}
    for fam, (prs, cfg) in vorlagen.items():
        for s, _pos, name in _titelnamen(prs, cfg, vc.FAMILIE_ALLE_STRATEGIEN[fam]):
            aus_vorlage[s] = name
    raus = {}
    for _, r in sd.iterrows():
        s = str(r[sp_a]).strip()
        fam = "" if pd.isna(r[sp_f]) else str(r[sp_f]).strip()
        raus[s] = (fam, aus_vorlage.get(s) or clean_strategy_name(s))
    return raus


def _pruefe_familie(titel, tabelle=None):
    print("\n5. Titel -> Familie (_STRATEGIE_FAMILIE in chart_dynamik.py)")
    if tabelle is None:
        tabelle = cd._STRATEGIE_FAMILIE
    fehler = 0
    for s, (fam, name) in titel.items():
        text = WE_TITLE_FORMAT.format(name=name)
        ist = cd._familie_aus_text(text, tabelle)
        if (ist or "").lower() != fam.lower():
            fehler += 1
            print(f"   FEHLER — '{s}': Titel '{text}' ergibt Familie {ist!r}, "
                  f"Mapping sagt '{fam}'. Ringe liefen mit fremdem Format.")
    alle = " ".join(WE_TITLE_FORMAT.format(name=n).lower()
                    for _f, n in titel.values())
    tot = sorted(k for k in tabelle
                 if not re.search(r"\b" + re.escape(k) + r"\b", alle))
    for k in tot:
        fehler += 1
        print(f"   FEHLER — Eintrag '{k}' trifft keinen Titel mehr (tot). "
              f"Entfernen oder an den neuen Titel anpassen.")
    if not fehler:
        print(f"   OK — {len(titel)} Strategien richtig zugeordnet, "
              f"alle {len(tabelle)} Eintraege in Gebrauch")
    return fehler


# ─────────────────────────────────────────────────────────────────────────
# 6. Gebaute Broschueren
# ─────────────────────────────────────────────────────────────────────────
def _pruefe_broschuere(prs, cfg, familie, etikett):
    """Wertentwicklungs-Titel = Titelname der Struktur-Folie; Familie erkannt."""
    fehler = 0
    for k, block in enumerate(cfg["feste_bloecke"]):
        if "anlagevorschlag" not in block or "wertentwicklung" not in block:
            continue
        name = strategiename_aus_titel(prs.slides[block["anlagevorschlag"] - 1])
        we = prs.slides[block["wertentwicklung"] - 1].shapes.title
        ist = _flach(we.text_frame.text) if we is not None else None
        soll = WE_TITLE_FORMAT.format(name=name)
        if ist != soll:
            fehler += 1
            print(f"   FEHLER — {etikett} F{block['wertentwicklung']}: "
                  f"'{ist}', erwartet '{soll}'")
    ist_fam = cd._familie_aus_prs(prs)
    if (ist_fam or "").lower() != familie.lower():
        fehler += 1
        print(f"   FEHLER — {etikett}: Familie am Artefakt {ist_fam!r}, "
              f"erwartet '{familie}'")
    return fehler


def _pruefe_gebaut(ausgabe):
    print("\n6. Gebaute Broschueren: Wertentwicklungs-Titel folgt der "
          "Struktur-Folie")
    try:
        # Daten und Bau aus dem Export-Smoketest — derselbe Pfad wie die
        # Oberflaeche, kein Nachbau (Lehre vom 23.09.2026).
        from test_export_smoke import _daten, _portfolio, _bauen
    except ImportError as ex:
        print(f"   UEBERSPRUNGEN — {ex}")
        return 0, {}
    d = _daten()
    fehler = 0
    gebaut = {}
    for fam, strategien in vc.FAMILIE_ALLE_STRATEGIEN.items():
        portfolios = [_portfolio(s, d) for s in strategien if s in d["d2c"]]
        ziel, _gr, meldungen = _bauen(portfolios, fam, d, ausgabe, f"{fam}.pptx")
        for m in meldungen:
            fehler += 1
            print(f"   FEHLER — {fam}: Build-Meldung {m[:90]}")
        prs = Presentation(ziel)
        gebaut[fam] = ziel
        f = _pruefe_broschuere(prs, vc.VORLAGEN_FAMILIEN[fam][1], fam, fam)
        fehler += f
        if not f:
            print(f"   OK — {fam}: {len(portfolios)} Strategien, Titel paarweise "
                  f"gleich, Familie erkannt")
    return fehler, gebaut


# ─────────────────────────────────────────────────────────────────────────
# 7. Gegenproben
# ─────────────────────────────────────────────────────────────────────────
def _gegenproben(vorlagen, kasten, titel, gebaut):
    print("\n7. Gegenproben: absichtlich verletzen, muss anschlagen")
    fehler = 0

    # (a) Struktur-Titel ohne "Anlagestrategie " -> Schritt 1.
    prs, cfg = _vorlage(VERGLEICH_FAMILIE)
    pos = cfg["feste_bloecke"][0]["anlagevorschlag"]
    prs.slides[pos - 1].shapes.title.text_frame.text = "Konservativ"
    if _still(_pruefe_titel_vorhanden, {VERGLEICH_FAMILIE: (prs, cfg)}) == 0:
        fehler += 1
        print("   FEHLER — Titel ohne Praefix blieb unbemerkt")
    else:
        print("   OK — Titel ohne Praefix wird gemeldet")

    # (b) Kasten umbenannt, Vorlage nicht -> Schritte 2, 3 und 4.
    kaputt = dict(kasten)
    kaputt["cVV defensiv plus"] = "Defensiv Premium"
    treffer = [_still(f, vorlagen, kaputt) for f in
               (_pruefe_kasten, _pruefe_uebersicht, _pruefe_vergleich)]
    if 0 in treffer:
        fehler += 1
        print(f"   FEHLER — halbe Umbenennung nicht ueberall gemeldet {treffer}")
    else:
        print("   OK — halbe Umbenennung meldet Kasten, Tabelle und Chart")

    # (c) Eintrag der Familien-Tabelle fehlt -> Schritt 5 (falsche Familie).
    tab = {k: v for k, v in cd._STRATEGIE_FAMILIE.items()
           if k != "esg defensiv plus"}
    if _still(_pruefe_familie, titel, tab) == 0:
        fehler += 1
        print("   FEHLER — fehlender Tabelleneintrag blieb unbemerkt")
    else:
        print("   OK — fehlender Eintrag wird gemeldet ('ESG Defensiv Plus' "
              "fiele sonst auf CVV)")

    # (d) Toter Eintrag -> Schritt 5.
    tab = dict(cd._STRATEGIE_FAMILIE, **{"comdirect 30": "comdirect"})
    if _still(_pruefe_familie, titel, tab) == 0:
        fehler += 1
        print("   FEHLER — toter Tabelleneintrag blieb unbemerkt")
    else:
        print("   OK — toter Eintrag wird gemeldet")

    # (e) Der alte Zustand am Artefakt -> Schritt 6.
    if "ESG" in gebaut:
        prs = Presentation(gebaut["ESG"])      # frisch geladen, nicht geteilt
        cfg = vc.VORLAGEN_FAMILIEN["ESG"][1]
        pos = cfg["feste_bloecke"][1]["wertentwicklung"]
        prs.slides[pos - 1].shapes.title.text_frame.text = \
            WE_TITLE_FORMAT.format(name="ESG defensiv+")
        if _still(_pruefe_broschuere, prs, cfg, "ESG", "ESG") == 0:
            fehler += 1
            print("   FEHLER — 'ESG defensiv+' neben 'ESG Defensiv Plus' "
                  "blieb unbemerkt")
        else:
            print("   OK — der Zustand vor dem 24.09.2026 wird gemeldet")
    return fehler


def _still(funktion, *args):
    """Pruefung ohne Ausgabe — sonst stuenden in den Gegenproben Zeilen, die
    wie echte Fehler aussehen."""
    with contextlib.redirect_stdout(io.StringIO()):
        return funktion(*args)


# ─────────────────────────────────────────────────────────────────────────
def main():
    ausgabe = (sys.argv[1] if len(sys.argv) > 1
               else tempfile.mkdtemp(prefix="ffpb_namen_"))
    os.makedirs(ausgabe, exist_ok=True)

    sd = stamm.lade()
    if sd is None or sd.empty:
        print(f"FEHLER: {stamm.PFAD} nicht ladbar")
        return 1
    kasten = _kasten_namen(sd)
    vorlagen = {fam: _vorlage(fam) for fam in vc.FAMILIE_ALLE_STRATEGIEN}
    titel = _titel_je_strategie(vorlagen, sd)

    fehler = (_pruefe_titel_vorhanden(vorlagen)
              + _pruefe_kasten(vorlagen, kasten)
              + _pruefe_uebersicht(vorlagen, kasten)
              + _pruefe_vergleich(vorlagen, kasten)
              + _pruefe_familie(titel))
    f6, gebaut = _pruefe_gebaut(ausgabe)
    fehler += f6 + _gegenproben(vorlagen, kasten, titel, gebaut)

    print()
    if fehler:
        print(f"FEHLGESCHLAGEN — {fehler} Abweichung(en)")
        return 1
    print("BESTANDEN — jede Strategie traegt auf Struktur-Folie, "
          "Wertentwicklungs-Folie, Kasten, Tabelle und Chart denselben Namen")
    return 0


if __name__ == "__main__":
    sys.exit(main())
