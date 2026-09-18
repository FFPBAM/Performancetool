"""Prueft die Inhaltsverzeichnisse aller Vorlagen (18.09.2026).

Die Seitenzahlen im Inhaltsverzeichnis sind in den Vorlagen EINGETIPPT. Nichts
im Code rechnet sie nach — ausser der PDF-Fassung, und die senkt nur um die
entfernte Vertriebsfolie (pdf_export._inhaltsverzeichnis_nachziehen). Eine
Zahl, die schon in der Vorlage falsch ist, bleibt es also in PowerPoint UND
im PDF.

AUSLOESER: Drei Eintraege zeigten auf die falsche Folie —
    ESG        Rechtliche Hinweise und Impressum  36 -> 37  (36 = "Unser Reporting")
    comdirect  Honorar                            14 -> 12  (14 = Honorar-Tabelle)
    comdirect  Rechtliche Hinweise und Impressum  25 -> 26  (25 = "Risikohinweise")
Zielfolien nach Entscheidung Philip: ESG/ETF zeigen auf die Trennfolie des
Abschnitts, cVV/comdirect auf den Abschnittsbeginn bzw. die gleichnamige
Folie. Davor (17.09.2026) schon cVV Impressum 34 -> 36.

Die Vorlage hat genau so viele Folien wie die gebaute Broschuere (gemessen
18.09.2026, alle Strategien) — die Nummer der Vorlage IST die Nummer im
Ergebnis. Schritt 1 haelt diese Voraussetzung fest.

    Schritt 1 — Folienzahl je Vorlage wie erwartet
    Schritt 2 — jeder Eintrag zeigt NAMENTLICH auf seine Zielfolie
    Schritt 3 — Gegenprobe: der alte ESG-Fehler (36) wird erkannt

    python tests/test_inhaltsverzeichnis.py
"""

import os
import re
import sys

WURZEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INHALT_LAYOUT = "Inhaltsverzeichnis"

# Vorlage -> (Folienzahl, [(Eintrag, Seitenzahl, Anfang des Zielfolien-Titels)])
ERWARTET = {
    "Vorlage_cVV_Infoboard.pptx": (37, [
        ("Unsere Vermögensverwaltung", 3, "Unsere Vermögensverwaltung"),
        ("Unsere klassischen Vermögensverwaltungsstrategien", 6,
         "Unsere fünf klassischen"),
        ("Unsere Strategien in Krisenphasen", 18, "Portfoliostresstest"),
        ("Wertentwicklung im historischen Verlauf", 19, "Die Wertentwicklung unserer"),
        ("Unsere individuelle Vermögensverwaltung", 20,
         "Unsere individuelle Vermögensverwaltung"),
        ("Honorar", 21, "Unser Honorar"),
        ("Unsere Investmentphilosophie", 24, "Unsere Investmentphilosophie"),
        ("Die Bank in Zahlen", 29, "Unsere Bank in Zahlen"),
        ("Unsere Standorte", 31, "Unsere Standorte"),
        ("Unsere Tradition", 33, "Unsere Tradition"),
        ("Rechtliche Hinweise und Impressum", 36, "Rechtliche Hinweise und Impressum"),
    ]),
    "Vorlage_ESG.pptx": (39, [
        ("ESG Basisinformationen", 4, "ESG BASIS-INFORMATIONEN"),
        ("Unsere ESG Strategien", 15, "UNSERE ESG-STRATEGIEN"),
        ("Fürst Fugger Privatbank und ihre Standorte", 25, "DIE FÜRST FUGGER"),
        ("Anlagen", 28, "ANLAGE"),
        ("Rechtliche Hinweise und Impressum", 37, "RECHTLICHE HINWEISE UND IMPRESSUM"),
    ]),
    "Vorlage_ETF.pptx": (35, [
        ("ESG Basisinformationen", 3, "ESG BASIS-INFORMATIONEN"),
        ("Unsere ESG-ETF Strategien", 15, "UNSERE ESG-ETF STRATEGIEN"),
        ("Fürst Fugger Privatbank und ihre Standorte", 21, "FÜRST FUGGER"),
        ("Anlagen", 24, "ANLAGE"),
        ("Rechtliche Hinweise und Impressum", 32, "RECHTLICHE HINWEISE UND IMPRESSUM"),
    ]),
    "Vorlage_comdirect.pptx": (27, [
        # Folie 3 heisst in der Vorlage "Unsere Vermögensverwaltung" — so steht
        # es dort, der Abschnitt beginnt trotzdem hier.
        ("Unsere Portfolioverwaltung", 3, "Unsere Vermögensverwaltung"),
        ("Unsere klassischen Portfolioverwaltungsstrategien", 5,
         "Unsere drei klassischen"),
        # 12 = Bildfolie "Unser Honorar" (Abschnittsbeginn), NICHT 14 (Tabelle,
        # traegt denselben Titel) — deshalb zaehlt hier die Nummer, nicht nur
        # der Titel.
        ("Honorar", 12, "Unser Honorar"),
        ("Unsere Investmentphilosophie", 16, "Unsere Investmentphilosophie"),
        ("Die Bank in Zahlen", 21, "Unsere Bank in Zahlen"),
        ("Unsere Tradition", 22, "Unsere Tradition"),
        ("Rechtliche Hinweise und Impressum", 26, "Rechtliche Hinweise und Impressum"),
    ]),
}


def _norm(text):
    return re.sub(r"\s+", " ", text.replace("\x0b", " ")).strip().casefold()


def _titel(folie):
    """Titel der Folie: Titel-Platzhalter zuerst, sonst der oberste Text."""
    kandidaten = []
    for sh in folie.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            ist_titel = sh.is_placeholder and str(
                sh.placeholder_format.type).startswith(("TITLE", "CENTER_TITLE"))
            kandidaten.append((0 if ist_titel else 1, sh.top, sh.text_frame.text))
    kandidaten.sort(key=lambda k: (k[0], k[1]))
    return kandidaten[0][2] if kandidaten else ""


def eintraege(prs):
    """[(Eintrag, Seitenzahl)] aller Inhaltsverzeichnis-Folien."""
    liste = []
    for folie in prs.slides:
        if folie.slide_layout.name != INHALT_LAYOUT:
            continue
        for sh in folie.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                text = "".join(r.text for r in p.runs)
                treffer = re.search(r"(\d+)\s*$", text)
                if "\t" in text and treffer:
                    liste.append((text.split("\t")[0].strip(), int(treffer.group(1))))
    return liste


def pruefe(prs, erwartet):
    """Liste der Abweichungen (leer = alles richtig)."""
    fehler = []
    ist = eintraege(prs)
    if [e for e, _ in ist] != [e for e, _, _ in erwartet]:
        fehler.append(f"Eintraege {[e for e, _ in ist]} statt "
                      f"{[e for e, _, _ in erwartet]}")
        return fehler
    for (eintrag, nr), (_, nr_soll, titel_soll) in zip(ist, erwartet):
        if nr != nr_soll:
            ziel = _titel(prs.slides[nr - 1]) if 0 < nr <= len(prs.slides) else "(fehlt)"
            fehler.append(f"'{eintrag}' zeigt auf {nr} ({ziel.strip()[:40]!r}) "
                          f"statt {nr_soll}")
        elif not _norm(_titel(prs.slides[nr - 1])).startswith(_norm(titel_soll)):
            fehler.append(f"'{eintrag}' -> {nr}: Folie heisst "
                          f"{_titel(prs.slides[nr - 1]).strip()[:40]!r}, "
                          f"erwartet '{titel_soll}'")
    return fehler


def main():
    from importlib.util import find_spec
    if find_spec("pptx") is None:
        print("UEBERSPRUNGEN — python-pptx nicht installiert")
        return 0
    from pptx import Presentation

    f = 0
    print("Schritt 1 + 2 — Folienzahl und Zielfolie jedes Eintrags")
    for vorlage, (anzahl, erwartet) in ERWARTET.items():
        prs = Presentation(os.path.join(WURZEL, "Vorlage", vorlage))
        if len(prs.slides) != anzahl:
            print(f"   FEHLER — {vorlage}: {len(prs.slides)} Folien statt {anzahl}; "
                  "alle Seitenzahlen im Inhaltsverzeichnis pruefen")
            f += 1
            continue
        abweichungen = pruefe(prs, erwartet)
        for a in abweichungen:
            print(f"   FEHLER — {vorlage}: {a}")
        f += len(abweichungen)
        if not abweichungen:
            print(f"   OK — {vorlage}: {len(erwartet)} Eintraege auf ihren Zielfolien")

    print("Schritt 3 — Gegenprobe: der alte ESG-Fehler (36) wird erkannt")
    prs = Presentation(os.path.join(WURZEL, "Vorlage", "Vorlage_ESG.pptx"))
    for folie in prs.slides:
        if folie.slide_layout.name != INHALT_LAYOUT:
            continue
        for sh in folie.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.endswith("Impressum\t37"):
                            r.text = r.text[:-2] + "36"
    if pruefe(prs, ERWARTET["Vorlage_ESG.pptx"][1]):
        print("   OK — die alte Zahl wuerde gemeldet")
    else:
        print("   FEHLER — die alte Zahl ginge durch, Schritt 2 prueft nichts")
        f += 1

    print()
    if f:
        print(f"FEHLGESCHLAGEN — {f} Abweichung(en)")
        return 1
    print("BESTANDEN — alle Inhaltsverzeichnisse zeigen auf ihre Zielfolien")
    return 0


if __name__ == "__main__":
    sys.exit(main())
