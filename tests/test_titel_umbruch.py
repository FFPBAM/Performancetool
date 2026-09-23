"""Prueft den Titel der rollierenden Folie der Thema-Broschueren (18.09.2026).

GEMELDET (Philip): Bei Pro Dividende (Folie 13) und den beiden SCHWEIZ-
Strategien (Folie 11) laeuft "Wertentwicklung der Strategie {Name}" ins Foto
rechts — in der PowerPoint und im PDF.

URSACHE: Der Master der Thema-Vorlagen stellt Titel auf wrap="none". Ein
langer Name laeuft nach rechts aus dem Titelfeld (endet 13,29 cm) ins Bild
(beginnt 13,36 cm). In PowerPoint gemessen, wo der Titel endet:

    Pro                          11,77 cm
    Offensiv                     13,27 cm   <- laengster Name, der passt
    Pro Dividende                15,03 cm
    Schweiz aktienorientiert     18,24 cm
    Schweiz substanzorientiert   18,99 cm

Alle anderen Titel der fuenf Thema-Broschueren liegen in ihrem Feld (ebenfalls
gemessen) — betroffen ist nur diese eine Folie.

REGEL (Entscheidung Philip): Umbruch nach "Strategie", aber nur wo noetig.
Namen bis ROLLIEREND_TITEL_EINZEILIG_MAX Zeichen bleiben einzeilig.

    Schritt 1 — jede Thema-Strategie aus Mapping_Namen.xlsx ist hier
                NAMENTLICH eingeordnet. Eine neue Strategie schlaegt an: vorher
                in PowerPoint messen, ob sie in eine Zeile passt.
    Schritt 2 — an allen drei Thema-Vorlagen: lange Namen ergeben genau zwei
                Zeilen in einem Absatz, kurze bleiben unveraendert.
    Schritt 3 — Gegenprobe: die alte Fassung (eine Zeile) wird erkannt.

    python tests/test_titel_umbruch.py
"""

import os
import sys

WURZEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, WURZEL)

ZEILE1 = "Wertentwicklung der Strategie"

# Namentlich: True = zweizeilig. Gemessen 18.09.2026 (siehe oben).
ERWARTET = {
    "Pro": False,
    "Offensiv": False,
    "Pro Dividende": True,
    "Schweiz aktienorientiert": True,
    "Schweiz substanzorientiert": True,
}

THEMA_VORLAGEN = [
    "Vorlage_Thema.pptx",
    "Vorlage_Thema_Offensiv.pptx",
    "Vorlage_Thema_ProDividende.pptx",
]


def _titel_der_rollierenden_folie(prs):
    """(Folienindex, Titel-Shape) der rollierenden Folie — erkannt am Titel."""
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if (sh.name in ("Titel", "Titel 2") and sh.has_text_frame
                    and sh.text_frame.text.startswith(ZEILE1)):
                return i, sh
    return None, None


def _zeilen(shape):
    """Sichtbare Zeilen des Titels: Absaetze, innerhalb davon <a:br>."""
    from pptx.oxml.ns import qn
    zeilen = []
    for p in shape.text_frame.paragraphs:
        aktuell = ""
        for kind in p._p:
            if kind.tag == qn("a:r"):
                aktuell += kind.find(qn("a:t")).text or ""
            elif kind.tag == qn("a:br"):
                zeilen.append(aktuell)
                aktuell = ""
        zeilen.append(aktuell)
    while zeilen and not zeilen[-1].strip():
        zeilen.pop()
    return zeilen


def _pruefe(shape, name):
    """Gibt eine Fehlerbeschreibung zurueck oder None."""
    zeilen = _zeilen(shape)
    absaetze = len(shape.text_frame.paragraphs)
    if ERWARTET[name]:
        if zeilen != [ZEILE1, name] or absaetze != 1:
            return f"erwartet 2 Zeilen in 1 Absatz, ist {zeilen} ({absaetze} Absaetze)"
    elif zeilen != [f"{ZEILE1} {name}"]:
        return f"erwartet 1 Zeile, ist {zeilen}"
    return None


def schritt1_namen_eingeordnet():
    print("Schritt 1 — jede Thema-Strategie ist namentlich eingeordnet")
    from modules.pptx_slides import (clean_strategy_name,
                                     ROLLIEREND_TITEL_EINZEILIG_MAX)
    from modules import stammdaten as stamm
    nm = stamm.lade()
    thema = [clean_strategy_name(str(n)) for n, f in
             zip(nm[stamm.spalte(nm, stamm.SP_ANZEIGE)],
                 nm[stamm.spalte(nm, stamm.SP_FAMILIE)])
             if str(f).strip() == "Thema"]
    f = 0
    for name in thema:
        if name not in ERWARTET:
            print(f"   FEHLER — '{name}' ist neu: in PowerPoint messen, ob der "
                  "Titel in eine Zeile passt, dann in ERWARTET eintragen")
            f += 1
        elif ERWARTET[name] != (len(name) > ROLLIEREND_TITEL_EINZEILIG_MAX):
            print(f"   FEHLER — '{name}' ({len(name)} Zeichen) passt nicht zur "
                  f"Grenze {ROLLIEREND_TITEL_EINZEILIG_MAX}")
            f += 1
    fehlend = sorted(set(ERWARTET) - set(thema))
    if fehlend:
        print(f"   HINWEIS — nicht mehr im Mapping: {fehlend}")
    if not f:
        print(f"   OK — {len(thema)} Strategien, Grenze "
              f"{ROLLIEREND_TITEL_EINZEILIG_MAX} Zeichen trennt sie wie gemessen")
    return f


def schritt2_vorlagen():
    print("Schritt 2 — Titel an allen Thema-Vorlagen")
    from pptx import Presentation
    from modules.pptx_slides import fill_rollierend_slide
    f = 0
    for vorlage in THEMA_VORLAGEN:
        pfad = os.path.join(WURZEL, "Vorlage", vorlage)
        for name in ERWARTET:
            prs = Presentation(pfad)
            idx, _ = _titel_der_rollierenden_folie(prs)
            if idx is None:
                print(f"   FEHLER — {vorlage}: keine rollierende Folie gefunden")
                f += 1
                break
            fill_rollierend_slide(prs, idx, name)
            _, titel = _titel_der_rollierenden_folie(prs)
            fehler = _pruefe(titel, name)
            if fehler:
                print(f"   FEHLER — {vorlage} / {name}: {fehler}")
                f += 1
            elif ERWARTET[name]:
                # Beide Zeilen tragen dieselbe Formatierung.
                from pptx.oxml.ns import qn
                rprs = [r.find(qn("a:rPr")) for r in
                        titel.text_frame.paragraphs[0]._p.findall(qn("a:r"))]
                from lxml import etree
                if len({etree.tostring(x) if x is not None else b"" for x in rprs}) != 1:
                    print(f"   FEHLER — {vorlage} / {name}: Zeilen unterschiedlich formatiert")
                    f += 1
    if not f:
        print(f"   OK — {len(THEMA_VORLAGEN)} Vorlagen x {len(ERWARTET)} Namen "
              "wie erwartet umbrochen")
    return f


def schritt3_gegenprobe():
    print("Schritt 3 — Gegenprobe: die alte einzeilige Fassung wird erkannt")
    from pptx import Presentation
    from modules.pptx_helpers import replace_text_in_shape
    prs = Presentation(os.path.join(WURZEL, "Vorlage", "Vorlage_Thema_ProDividende.pptx"))
    _, titel = _titel_der_rollierenden_folie(prs)
    replace_text_in_shape(titel, f"{ZEILE1} Pro Dividende")
    if _pruefe(titel, "Pro Dividende") is None:
        print("   FEHLER — die alte Fassung ginge durch, Schritt 2 prueft nichts")
        return 1
    print("   OK — die einzeilige Fassung wuerde gemeldet")
    return 0


def main():
    from importlib.util import find_spec
    fehlt = [p for p in ("pptx", "pandas") if find_spec(p) is None]
    if fehlt:
        print(f"UEBERSPRUNGEN — nicht installiert: {fehlt}")
        return 0
    fehler = schritt1_namen_eingeordnet() + schritt2_vorlagen() + schritt3_gegenprobe()
    print()
    if fehler:
        print(f"FEHLGESCHLAGEN — {fehler} Abweichung(en)")
        return 1
    print("BESTANDEN — Titel der rollierenden Folie bleibt im Titelfeld")
    return 0


if __name__ == "__main__":
    sys.exit(main())
