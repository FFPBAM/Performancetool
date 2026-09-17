# modules/pdf_export.py
"""
PDF-Fassung der Broschuere — die PowerPoint-Quelle dafuer (NEU 17.09.2026).

WARUM ES DIESES MODUL GIBT: Neben der PowerPoint soll die Broschuere auch als
PDF herausgehen. Das PDF entsteht aus DERSELBEN gebauten PowerPoint (dieselbe
Befuellung, kein zweiter Bauweg) — nur ohne die Folie "Ihre Ansprechpartner
fuer den Vertrieb". Grund (Philip): In der PowerPoint tauscht der Berater die
Fotos und Namen gegen seine eigenen aus; im PDF geht das nicht, also gehoert
die Folie dort nicht hinein.

Dieses Modul baut die PDF-FASSUNG: eine PowerPoint, in der
  1. jede Folie mit dem Layout "Ansprechpartner" entfernt ist,
  2. die Seitenzahlen unten neu geschrieben sind,
  3. die festen Seitenzahlen im Inhaltsverzeichnis nachgezogen sind,
  4. Diagramm-Verknuepfungen auf externe Mappen (Netzlaufwerk) entfernt sind.
Die Umwandlung in PDF macht der Berater in PowerPoint ("Speichern unter ->
PDF"). Ein automatischer Umwandlungsweg ist ohne Freigabe der IT-Sicherheit
nicht vorgesehen (Stand 17.09.2026).

WARUM AM LAYOUT UND NICHT AN DER FOLIENNUMMER: Die Folie steckt heute nur in
der cVV-Vorlage (F30) und in der ungenutzten Standard-Vorlage (F21) — beide
mit dem Layout "Ansprechpartner". Eine feste Nummer waere bei der naechsten
Vorlagenaenderung falsch; das Layout wandert mit der Folie. Bekommt eine
weitere Familie die Folie, greift es ohne Codeaenderung.

WARUM DAS INHALTSVERZEICHNIS NACHGEZOGEN WERDEN MUSS: Es traegt keine Felder,
sondern eingetippte Zahlen ("Unsere Standorte <Tab> 31"). Faellt eine Folie
davor weg, zeigen sie ins Leere. `update_slide_numbers` hilft hier nicht —
es beschreibt nur die Seitenzahl-Shapes unten auf der Folie.

Streamlit-frei wie `pptx_export` (Batch-Faehigkeit).
"""

import io
import posixpath
import re
import zipfile

from lxml import etree
from pptx import Presentation

try:
    from modules.pptx_helpers import remove_slide, update_slide_numbers
except ImportError:
    from pptx_helpers import remove_slide, update_slide_numbers


VERTRIEB_LAYOUT = "Ansprechpartner"
"""Name des Folienlayouts, dessen Folien NICHT ins PDF gehoeren.

Als Name im Layout (`<p:cSld name="...">`), nicht als Folientitel: Der Titel
lautet "Ihre Ansprechpartner<Zeilenumbruch>fuer den Vertrieb" und ist Text,
den jemand in der Vorlage umformulieren kann. Der Layoutname ist Struktur."""

INHALT_LAYOUT = "Inhaltsverzeichnis"
"""Layoutname der Inhaltsverzeichnis-Folie (cVV, ESG, ETF, comdirect, FFPB)."""

HINWEIS_VERTRIEB = (
    "Lädt eine PowerPoint für das PDF: ohne die Folie „Ihre Ansprechpartner "
    "für den Vertrieb“ (im PDF lassen sich die Bilder nicht austauschen), "
    "Seitenzahlen und Inhaltsverzeichnis angepasst. In PowerPoint über "
    "„Speichern unter → PDF“ sichern.")
"""Tooltip am PDF-Button, wenn die Vorlage der Familie die Folie fuehrt."""

HINWEIS_OHNE_VERTRIEB = (
    "Lädt dieselbe Broschüre als PowerPoint für das PDF. In PowerPoint über "
    "„Speichern unter → PDF“ sichern.")
"""Tooltip am PDF-Button, wenn nichts entfernt wird."""


_NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_TYP_LAYOUT = _NS_R + "/slideLayout"

_ZAHL_AM_ENDE = re.compile(r"(\d+)(\s*)$")


def _rels(zf, teil):
    """Beziehungen eines Teils als {rId: (Typ, Zielteil)}."""
    ordner, datei = posixpath.split(teil)
    pfad = posixpath.join(ordner, "_rels", datei + ".rels")
    if pfad not in zf.namelist():
        return {}
    wurzel = etree.fromstring(zf.read(pfad))
    raus = {}
    for rel in wurzel.iter("{%s}Relationship" % _NS_REL):
        ziel = rel.get("Target")
        if rel.get("TargetMode") != "External":
            ziel = posixpath.normpath(posixpath.join(ordner, ziel))
        raus[rel.get("Id")] = (rel.get("Type"), ziel)
    return raus


def folien_layouts(pptx_quelle) -> list:
    """Layoutname je Folie, in Folienreihenfolge.

    Liest direkt am ZIP (kein python-pptx), damit es auch fuer die 17-MB-
    Vorlagen schnell genug ist, um beim Zeichnen der Oberflaeche zu laufen.

    Args:
        pptx_quelle: Dateipfad oder PPTX-Bytes.
    """
    if isinstance(pptx_quelle, (bytes, bytearray)):
        pptx_quelle = io.BytesIO(pptx_quelle)
    with zipfile.ZipFile(pptx_quelle) as zf:
        praes = etree.fromstring(zf.read("ppt/presentation.xml"))
        praes_rels = _rels(zf, "ppt/presentation.xml")
        liste = praes.find("{%s}sldIdLst" % _NS_P)
        raus = []
        for sld in (liste if liste is not None else []):
            folie = praes_rels[sld.get("{%s}id" % _NS_R)][1]
            layout = next((ziel for typ, ziel in _rels(zf, folie).values()
                           if typ == _TYP_LAYOUT), None)
            name = None
            if layout:
                csld = etree.fromstring(zf.read(layout)).find(
                    "{%s}cSld" % _NS_P)
                name = csld.get("name") if csld is not None else None
            raus.append(name)
        return raus


def vertriebsfolien(pptx_quelle) -> list:
    """1-basierte Positionen der Folien, die nicht ins PDF gehoeren."""
    return [i for i, name in enumerate(folien_layouts(pptx_quelle), start=1)
            if name == VERTRIEB_LAYOUT]


def _inhaltsverzeichnis_nachziehen(prs, entfernt: list, meldungen: list):
    """Senkt jede eingetippte Seitenzahl um die Zahl der davor entfernten Folien.

    Die Zahl steht am Absatzende — bei cVV/comdirect als eigener Textlauf
    ("31"), bei ESG/ETF im selben Lauf wie der Eintrag ("Anlagen<Tab>28").
    Deshalb wird im LETZTEN Lauf mit einer Endziffer nur diese Ziffernfolge
    ersetzt; Schrift und Formatierung bleiben unangetastet.

    Laeuft NACH dem Entfernen: `prs` ist bereits verkuerzt, die Zahlen im
    Text beziehen sich aber noch auf die urspruengliche Folge.
    """
    for folie in prs.slides:
        if folie.slide_layout.name != INHALT_LAYOUT:
            continue
        for shape in folie.shapes:
            if not shape.has_text_frame:
                continue
            for absatz in shape.text_frame.paragraphs:
                if "\t" not in "".join(r.text for r in absatz.runs):
                    continue
                for lauf in reversed(absatz.runs):
                    treffer = _ZAHL_AM_ENDE.search(lauf.text)
                    if not treffer:
                        if lauf.text.strip():
                            break
                        continue
                    alt = int(treffer.group(1))
                    if alt in entfernt:
                        meldungen.append(
                            f"Inhaltsverzeichnis verweist auf Folie {alt}, "
                            f"die im PDF fehlt — Eintrag bitte prüfen.")
                        break
                    neu = alt - sum(1 for p in entfernt if p < alt)
                    if neu != alt:
                        lauf.text = (lauf.text[:treffer.start(1)] + str(neu)
                                     + treffer.group(2))
                    break


_TYP_HYPERLINK = _NS_R + "/hyperlink"
_NS_C = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def externe_verknuepfungen(pptx_quelle) -> list:
    """Alle externen Beziehungen AUSSER Weblinks: [(Teil, Typ-Endung)].

    WARUM (Sicherheitspruefung 17.09.2026): Diagramme koennen auf externe
    Mappen verweisen (`c:externalData` -> oleObject, TargetMode External). Fuer
    die Anzeige sind solche Verweise ohne Belang, in einer Datei, die
    weitergegeben wird, haben sie nichts zu suchen — sie verraten Pfade und
    laden beim Bearbeiten zu Verbindungen dorthin ein. Die PDF-Fassung wird
    deshalb davon befreit (`_externe_daten_entfernen`).
    """
    if isinstance(pptx_quelle, (bytes, bytearray)):
        pptx_quelle = io.BytesIO(pptx_quelle)
    raus = []
    with zipfile.ZipFile(pptx_quelle) as zf:
        for name in zf.namelist():
            if not name.endswith(".rels"):
                continue
            for rel in etree.fromstring(zf.read(name)).iter(
                    "{%s}Relationship" % _NS_REL):
                if (rel.get("TargetMode") == "External"
                        and rel.get("Type") != _TYP_HYPERLINK):
                    raus.append((name, rel.get("Type", "").rsplit("/", 1)[-1]))
    return raus


def _externe_daten_entfernen(prs) -> int:
    """Entfernt `c:externalData`, das auf eine EXTERNE Mappe zeigt, samt
    Beziehung. Die Diagramme zeichnen aus ihrem Zwischenspeicher (c:*Cache) —
    fuers PDF aendert sich nichts. Rueckgabe: Anzahl entfernter Verweise."""
    anzahl = 0
    for teil in prs.part.package.iter_parts():
        wurzel = getattr(teil, "_element", None)
        if wurzel is None or not str(teil.partname).startswith("/ppt/charts/chart"):
            continue
        for ext in wurzel.findall("{%s}externalData" % _NS_C):
            rid = ext.get("{%s}id" % _NS_R)
            rel = teil.rels.get(rid) if rid else None
            if rel is None or not rel.is_external:
                continue
            wurzel.remove(ext)
            teil.drop_rel(rid)
            anzahl += 1
    return anzahl


def pptx_fuer_pdf(pptx_bytes: bytes) -> tuple:
    """Baut aus der fertigen Broschuere die Quelle fuer das PDF.

    Returns:
        (bytes, entfernte_positionen, meldungen)
        - entfernte_positionen: 1-basiert, bezogen auf die Eingabe
        - meldungen: Auffaelligkeiten, die der Aufrufer anzeigen muss
          (gleiches Prinzip wie `pptx_export.LAST_BUILD_ERRORS`, nie still)
        Ohne Vertriebsfolie UND ohne externe Verknuepfungen kommen die
        Eingabe-Bytes unveraendert zurueck.
    """
    entfernt = vertriebsfolien(pptx_bytes)
    if not entfernt and not externe_verknuepfungen(pptx_bytes):
        return pptx_bytes, [], []

    meldungen = []
    prs = Presentation(io.BytesIO(pptx_bytes))
    # Von hinten entfernen, damit die vorderen Positionen gueltig bleiben.
    for pos in sorted(entfernt, reverse=True):
        remove_slide(prs, pos - 1)
    if entfernt:
        update_slide_numbers(prs)
        _inhaltsverzeichnis_nachziehen(prs, entfernt, meldungen)
    _externe_daten_entfernen(prs)

    puffer = io.BytesIO()
    prs.save(puffer)
    daten = puffer.getvalue()

    # Gegenkontrolle am Ergebnis statt Vertrauen in remove_slide.
    rest = vertriebsfolien(daten)
    if rest:
        meldungen.append(
            f"PDF-Quelle enthält weiter Vertriebsfolien an Position {rest}.")
    extern = externe_verknuepfungen(daten)
    if extern:
        meldungen.append(
            f"PDF-Fassung enthält weiter {len(extern)} externe Verknüpfung(en) "
            f"({extern[0][0]}) — bitte melden.")
    return daten, entfernt, meldungen
